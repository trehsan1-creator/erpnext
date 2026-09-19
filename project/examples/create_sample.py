"""Create a synthetic Iranian ledger input workbook."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

output = Path(__file__).with_name("sample_input.xlsx")
wb = Workbook(); ws = wb.active; ws.title = "اسناد خام"; ws.sheet_view.rightToLeft = True
headers = ["شماره سند", "تاریخ سند", "شرح", "بدهکار", "بستانکار", "کد حساب", "نام حساب", "عطف"]
rows = [
    [1001, "1403/01/15", "پرداخت اجاره دفتر فروردین", 85000000, 0, None, None, "BNK-101"],
    [1001, "۱۴۰۳/۰۱/۱۵", "برداشت از بانک بابت اجاره", 0, 85000000, "110101", "بانک‌ها", "BNK-101"],
    [1002, "1403/01/29", "فروش خدمات مشاوره مالی", 0, 125000000, None, None, "INV-22"],
    [1002, "1403/01/29", "مطالبات شرکت نمونه", 125000000, 0, "120101", "حساب‌های دریافتنی", "INV-22"],
    [1003, "1403/02/05", "واریزی بابت تسویه مورد قبلی", 42000000, 0, None, None, "TRX-778"],
    [1003, "1403/02/05", "طرف حساب نامشخص", 0, 42000000, None, None, "TRX-778"],
    [1004, "1403/02/30", "پرداخت حقوق اردیبهشت", 190000000, 0, None, None, "PAY-02"],
    [1004, "1403/02/30", "بانک بابت لیست حقوق", 0, 190000000, "110101", "بانک‌ها", "PAY-02"],
    [1005, "تاریخ خراب", "ردیف معیوب برای آزمون تحمل خطا", "نامعتبر", 0, None, None, None],
    [1006, "1403/03/10", "خرید ملزومات اداری", 15000000, 0, None, None, "BUY-8"],
]
ws.append(headers)
for row in rows: ws.append(row)
for cell in ws[1]:
    cell.font = Font(name="Vazirmatn", bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="17365D"); cell.alignment = Alignment(horizontal="center")
for col, width in zip("ABCDEFGH", [14, 16, 36, 18, 18, 15, 25, 15]): ws.column_dimensions[col].width = width
for row in range(2, ws.max_row + 1):
    ws.cell(row, 4).number_format = '#,##0'; ws.cell(row, 5).number_format = '#,##0'
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
wb.save(output); print(output)
