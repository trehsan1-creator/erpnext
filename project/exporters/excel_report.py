"""Professional Persian accounting workbook built directly with openpyxl."""
from __future__ import annotations
import json
from copy import copy
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from core.models import PipelineResult, ReviewStatus


class PersianExcelReport:
    MONEY_FORMAT = '#,##0;[Red](#,##0);-'

    def __init__(self, project_root: Path) -> None:
        style_path = project_root / "exporters" / "templates" / "default.json"
        self.theme: dict[str, str] = json.loads(style_path.read_text(encoding="utf-8"))
        self.font_name = self.theme["font"]

    def export(self, result: PipelineResult, output: Path) -> Path:
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        summary = workbook.active
        summary.title = "خلاصه مدیریتی"
        self._summary(summary, result)
        self._journal(workbook.create_sheet("دفتر روزنامه"), result)
        self._documents(workbook.create_sheet("کنترل اسناد"), result)
        self._warnings(workbook.create_sheet("هشدارها"), result)
        self._invalid(workbook.create_sheet("ردیف‌های نامعتبر"), result)
        for sheet in workbook.worksheets:
            self._finalize(sheet)
        workbook.save(output)
        return output

    def _summary(self, ws: Any, result: PipelineResult) -> None:
        ws.merge_cells("A1:F2")
        ws["A1"] = "داشبورد کنترل حسابداری"
        ws["A1"].font = Font(name=self.font_name, size=20, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor=self.theme["header_fill"])
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        lines = result.lines
        total_debit = sum((x.debit for x in lines), Decimal(0))
        total_credit = sum((x.credit for x in lines), Decimal(0))
        pending = sum(x.status != ReviewStatus.RESOLVED for x in lines)
        cards = [
            ("تعداد اسناد", len(result.documents)), ("تعداد ردیف‌ها", len(lines)),
            ("جمع بدهکار (ریال)", total_debit), ("جمع بستانکار (ریال)", total_credit),
            ("اختلاف تراز (ریال)", total_debit-total_credit), ("نیازمند بررسی", pending),
            ("هشدارها", len(result.warnings)), ("ردیف نامعتبر", len(result.invalid_rows)),
        ]
        for idx, (label, value) in enumerate(cards):
            row = 4 + (idx // 2) * 2
            col = 1 + (idx % 2) * 3
            ws.cell(row, col, label)
            ws.cell(row, col).font = Font(name=self.font_name, bold=True, color="FFFFFF")
            ws.cell(row, col).fill = PatternFill("solid", fgColor=self.theme["header_fill"])
            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+1)
            ws.cell(row+1, col, float(value) if isinstance(value, Decimal) else value)
            ws.cell(row+1, col).font = Font(name=self.font_name, size=15, bold=True, color=self.theme["header_fill"])
            ws.cell(row+1, col).fill = PatternFill("solid", fgColor=self.theme["accent_fill"])
            ws.cell(row+1, col).number_format = self.MONEY_FORMAT
            ws.merge_cells(start_row=row+1, start_column=col, end_row=row+1, end_column=col+1)
        ws["A14"] = "وضعیت"
        ws["B14"] = "تعداد"
        status_data = [("حل‌شده", sum(x.status == ReviewStatus.RESOLVED for x in lines)),
                       ("در انتظار AI", sum(x.status == ReviewStatus.PENDING_AI for x in lines)),
                       ("بررسی انسانی", sum(x.status == ReviewStatus.HUMAN_REVIEW for x in lines))]
        for row, item in enumerate(status_data, 15):
            ws.cell(row, 1, item[0]); ws.cell(row, 2, item[1])
        self._style_header(ws, 14, 2)
        chart = BarChart(); chart.title = "وضعیت پردازش"; chart.legend = None
        chart.add_data(Reference(ws, min_col=2, min_row=14, max_row=17), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=15, max_row=17)); chart.height = 6; chart.width = 12
        ws.add_chart(chart, "D14")
        ws["A20"] = "تولید گزارش"
        ws["B20"] = result.generated_at.astimezone().strftime("%Y-%m-%d %H:%M")
        ws["A21"] = "فایل منبع"; ws["B21"] = result.source_file
        ws.freeze_panes = "A4"

    def _journal(self, ws: Any, result: PipelineResult) -> None:
        headers = ["ردیف", "شماره سند", "تاریخ شمسی", "تاریخ میلادی", "شرح", "کد حساب", "نام حساب",
                   "بدهکار (ریال)", "بستانکار (ریال)", "وضعیت", "اطمینان", "یادداشت بررسی", "ردیف منبع"]
        ws.append(headers)
        status_names = {"resolved": "حل‌شده", "pending_ai": "در انتظار AI", "human_review": "بررسی انسانی", "invalid": "نامعتبر"}
        for idx, line in enumerate(result.lines, 1):
            ws.append([idx, line.document_number, line.jalali_date, line.gregorian_date, line.description,
                line.account_code, line.account_name, float(line.debit), float(line.credit), status_names[line.status.value],
                line.confidence, line.review_note, line.source_row])
        self._style_header(ws, 1, len(headers)); ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        ws.column_dimensions["D"].hidden = True
        for row in range(2, ws.max_row + 1):
            ws.cell(row, 4).number_format = "yyyy-mm-dd"
            for col in (8, 9): ws.cell(row, col).number_format = self.MONEY_FORMAT
            ws.cell(row, 11).number_format = "0%"
        if ws.max_row > 1:
            table = Table(displayName="JournalTable", ref=f"A1:M{ws.max_row}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
            ws.add_table(table)
            yellow = PatternFill("solid", fgColor=self.theme["warning_fill"])
            ws.conditional_formatting.add(f"A2:M{ws.max_row}", FormulaRule(formula=['$J2<>"حل‌شده"'], fill=yellow))

    def _documents(self, ws: Any, result: PipelineResult) -> None:
        headers = ["شماره سند", "تاریخ شمسی", "تعداد ردیف", "جمع بدهکار", "جمع بستانکار", "اختلاف", "وضعیت تراز"]
        ws.append(headers)
        for doc in result.documents:
            ws.append([doc.number, doc.jalali_date, len(doc.lines), float(doc.total_debit), float(doc.total_credit),
                       float(doc.total_debit-doc.total_credit), "تراز" if doc.is_balanced else "نامتوازن"])
        self._style_header(ws, 1, len(headers)); ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for row in range(2, ws.max_row+1):
            for col in (4, 5, 6): ws.cell(row, col).number_format = self.MONEY_FORMAT
        if ws.max_row > 1:
            ws.conditional_formatting.add(f"A2:G{ws.max_row}", FormulaRule(formula=['$G2="نامتوازن"'], fill=PatternFill("solid", fgColor=self.theme["error_fill"])))

    def _warnings(self, ws: Any, result: PipelineResult) -> None:
        ws.append(["ردیف", "سطح", "کد", "پیام", "شماره سند", "ردیف منبع"])
        for idx, warning in enumerate(result.warnings, 1):
            ws.append([idx, warning.severity.value, warning.code, warning.message, warning.document_number, warning.row_number])
        if not result.warnings: ws.append([1, "info", "OK", "هشداری ثبت نشده است", None, None])
        self._style_header(ws, 1, 6); ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

    def _invalid(self, ws: Any, result: PipelineResult) -> None:
        ws.append(["ردیف منبع", "خطا", "داده خام"])
        for row in result.invalid_rows:
            ws.append([row["source_row"], row["error"], json.dumps(row["raw_data"], ensure_ascii=False, default=str)])
        if not result.invalid_rows: ws.append([None, "ردیف نامعتبری ثبت نشده است", None])
        self._style_header(ws, 1, 3); ws.freeze_panes = "A2"

    def _style_header(self, ws: Any, row: int, columns: int) -> None:
        for col in range(1, columns+1):
            cell = ws.cell(row, col); cell.font = Font(name=self.font_name, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=self.theme["header_fill"])
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 28

    def _finalize(self, ws: Any) -> None:
        ws.sheet_view.rightToLeft = True
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.outlinePr.summaryBelow = True
        ws.oddHeader.center.text = f"&B{ws.title}"
        ws.oddFooter.center.text = "صفحه &P از &N"; ws.oddFooter.right.text = "گزارش حسابداری محرمانه"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        thin = Side(style="thin", color=self.theme["border_color"])
        for row in ws.iter_rows():
            for cell in row:
                font = copy(cell.font); font.name = self.font_name; cell.font = font
                alignment = copy(cell.alignment)
                alignment.vertical = "center"; alignment.readingOrder = 2
                alignment.wrap_text = cell.column in (4, 5, 12); cell.alignment = alignment
                if cell.value is not None: cell.border = Border(bottom=thin)
        for idx, column in enumerate(ws.columns, 1):
            values = [str(c.value) if c.value is not None else "" for c in column]
            width = min(max(max((len(v) for v in values), default=8) + 3, 10), 55)
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.sheet_view.showGridLines = False
